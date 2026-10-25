"""PrismDeck optional private document/video converter.

Launch behind TLS with a unique PRISMDECK_INGEST_TOKEN. No public anonymous
conversion route; source data is deleted after each request. Requires
LibreOffice + ffmpeg in the container and Whisper tiny model for audio.
"""
import hmac
import os
import re
import subprocess
import tempfile
from pathlib import Path
from typing import Any

from fastapi import FastAPI, HTTPException, Request, UploadFile, File
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import JSONResponse

MAX_BYTES = 120 * 1024 * 1024
MAX_CHARS = 150000
ALLOWED = {".doc", ".ppt", ".odt", ".odp", ".ods", ".mp4", ".mov", ".m4v",
           ".webm", ".mkv", ".mp3", ".wav", ".m4a"}
VIDEO = {".mp4", ".mov", ".m4v", ".webm", ".mkv", ".mp3", ".wav", ".m4a"}
ORIGINS = [x.strip() for x in os.getenv("PRISMDECK_ALLOWED_ORIGINS",
 "http://localhost:5173,http://127.0.0.1:5173").split(",") if x.strip()]

app = FastAPI(title="PrismDeck private ingestion", docs_url=None, redoc_url=None)
app.add_middleware(CORSMiddleware, allow_origins=ORIGINS,
 allow_methods=["POST", "GET"], allow_headers=["Authorization", "Content-Type"],
 allow_credentials=False)

@app.middleware("http")
async def headers(request: Request, call_next):
    reply = await call_next(request)
    reply.headers["X-Content-Type-Options"] = "nosniff"
    reply.headers["Cache-Control"] = "no-store"
    return reply

@app.get("/api/health")
def health():
    return {"status": "ready", "anonymous_ingestion": False}

def valid_token(auth: str) -> bool:
    secret = os.environ.get("PRISMDECK_INGEST_TOKEN", "")
    if len(secret) < 24 or not auth.startswith("Bearer "):
        return False
    return hmac.compare_digest(auth[7:], secret)

def run(*args: str, timeout: int = 90) -> str:
    try:
        out = subprocess.run(args, stdout=subprocess.PIPE, stderr=subprocess.PIPE,
                             text=True, timeout=timeout, check=True)
        return out.stdout
    except (subprocess.TimeoutExpired, subprocess.CalledProcessError, OSError) as exc:
        raise HTTPException(422, "Conversion failed or timed out.") from exc

def from_pptx(path: Path) -> str:
    from pptx import Presentation
    p = Presentation(str(path))
    lines = []
    for i, slide in enumerate(p.slides[:120], 1):
        items = [shape.text.strip() for shape in slide.shapes if shape.has_text_frame and shape.text.strip()]
        lines.append(f"[Slide {i}] " + " ".join(items))
    return "\n\n".join(lines)[:MAX_CHARS]

def from_docx(path: Path) -> str:
    from docx import Document
    doc = Document(str(path))
    paragraphs = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
    for t in doc.tables:
        for row in t.rows:
            paragraphs.append(" | ".join(cell.text.strip() for cell in row.cells))
    return "\n".join(paragraphs)[:MAX_CHARS]

def from_xlsx(path: Path) -> tuple[str, list[dict[str, Any]]]:
    from openpyxl import load_workbook
    wb = load_workbook(path, read_only=True, data_only=True)
    lines, series = [], []
    for sheet in wb.worksheets[:8]:
        lines.append("Sheet: " + sheet.title)
        for row in sheet.iter_rows(min_row=1, max_row=min(sheet.max_row or 1, 250),
                                   max_col=min(sheet.max_column or 1, 16), values_only=True):
            lines.append(" | ".join("" if x is None else str(x)[:180] for x in row))
            if len(row) >= 2 and isinstance(row[0], str) and isinstance(row[1], (int, float)) and len(series) < 8:
                series.append({"label": row[0][:35], "value": row[1]})
            if sum(map(len, lines)) > MAX_CHARS:
                break
    wb.close()
    return "\n".join(lines)[:MAX_CHARS], series

def office_extract(path: Path, ext: str, folder: Path):
    desired = ".docx" if ext in {".doc", ".odt"} else ".pptx" if ext in {".ppt", ".odp"} else ".xlsx"
    format_name = {"docx": "docx", "pptx": "pptx", "xlsx": "xlsx"}[desired.lstrip(".")]
    run("soffice", "-env:UserInstallation=file://" + str(folder / "lo-profile"),
        "--headless", "--convert-to", format_name, "--outdir", str(folder),
        str(path), timeout=70)
    converted = folder / (path.stem + desired)
    if not converted.is_file():
        raise HTTPException(422, "Legacy Office conversion produced no supported document.")
    if desired == ".pptx":
        return {"text": from_pptx(converted), "kind": "presentation", "numericSeries": []}
    if desired == ".docx":
        return {"text": from_docx(converted), "kind": "word", "numericSeries": []}
    text, series = from_xlsx(converted)
    return {"text": text, "kind": "spreadsheet", "numericSeries": series}

def video_extract(path: Path, folder: Path):
    duration_text = run("ffprobe", "-v", "error", "-show_entries",
       "format=duration", "-of", "default=noprint_wrappers=1:nokey=1",
       str(path), timeout=25).strip()
    try:
        duration = float(duration_text)
    except ValueError as exc:
        raise HTTPException(422, "Unable to determine video or audio duration.") from exc
    if not (0 < duration <= 900):
        raise HTTPException(422, "Transcription accepts media up to 15 minutes.")
    out = folder / "transcription.wav"
    run("ffmpeg", "-nostdin", "-hide_banner", "-loglevel", "error",
        "-i", str(path), "-vn", "-ac", "1", "-ar", "16000", "-y", str(out), timeout=120)
    # Lazy import prevents expensive model initialization on server startup.
    from faster_whisper import WhisperModel
    model = WhisperModel(os.getenv("PRISMDECK_WHISPER_MODEL", "tiny"),
                         device="cpu", compute_type="int8")
    segments, _ = model.transcribe(str(out), beam_size=2, vad_filter=True,
                                    condition_on_previous_text=False)
    text = " ".join(segment.text.strip() for segment in segments).strip()
    if not text:
        raise HTTPException(422, "No speech detected. Upload subtitles or a written transcript.")
    return {"text": text[:MAX_CHARS], "kind": "transcript", "numericSeries": [],
            "note": "Automated speech transcription, not human-verified. Review names, figures and quotations."}

@app.post("/api/ingest")
async def ingest(request: Request, file: UploadFile = File(...)):
    if not valid_token(request.headers.get("authorization", "")):
        raise HTTPException(401, "Authentication required.")
    name = Path(file.filename or "source.bin").name[:130]
    ext = Path(name).suffix.lower()
    if ext not in ALLOWED:
        raise HTTPException(415, "Unsupported converter input.")
    with tempfile.TemporaryDirectory(prefix="prismdeck_") as scratch:
        folder = Path(scratch)
        target = folder / ("source" + ext)
        total = 0
        with target.open("wb") as handle:
            while True:
                part = await file.read(1024 * 1024)
                if not part:
                    break
                total += len(part)
                if total > MAX_BYTES:
                    raise HTTPException(413, "File exceeds 120 MB.")
                handle.write(part)
        if total == 0:
            raise HTTPException(422, "Empty file.")
        result = video_extract(target, folder) if ext in VIDEO else office_extract(target, ext, folder)
        if not result["text"].strip():
            raise HTTPException(422, "No readable content extracted.")
        return JSONResponse(result)
